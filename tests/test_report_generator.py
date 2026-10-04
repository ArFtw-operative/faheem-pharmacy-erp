"""SQL-backed report generation: disclosure boundaries, dates, totals, stock and exports."""
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

import pytest
from sqlalchemy import select

from app.models import Customer, Item, Purchase, PurchaseItem, Supplier, Sale, SalePayment, InventoryMovement, User, Role, Permission
from app.services import inventory_service as inv, sales_service, report_generator as reports
from tests.conftest import login


def stocked(db,name='Report item',category='Medicine',qty=20):
    item=inv.create_item(db,name=name,category=category,manufacturer='Test Brand',base_unit='TABLET',pack_unit='TABLET',units_per_pack=1)
    batch=inv.add_or_update_batch(db,item,batch_no=name,expiry_date=date(2099,1,1),quantity=qty,purchase_rate='5',mrp='10',movement_type='OPENING_STOCK')
    db.commit();return item,batch


def sell(db,item,batch,qty=1,when=datetime(2026,9,1,12),**kwargs):
    sale=sales_service.create_sale(db,lines=[{'item_id':item.id,'batch_id':batch.id,'quantity':qty}],**kwargs)
    sale.sale_date=when;db.commit();return sale


RANGE={'period':'custom','from':'2026-09-01','to':'2026-09-30'}


def test_supplier_invoice_audit_selection_and_partial_receipt(db,client):
    supplier=Supplier(name='Audit Supplier');other=Supplier(name='Other supplier')
    db.add_all([supplier,other]);db.flush()
    docs=[]
    for sup,number in ((supplier,'A1'),(supplier,'A2'),(other,'B1')):
        doc=Purchase(supplier_id=sup.id,invoice_no=number,status='PARTIAL',total=Decimal('105'),
                     supplier_total=Decimal('106'),purchase_date=datetime(2026,9,2,12))
        db.add(doc);db.flush();docs.append(doc)
        db.add(PurchaseItem(purchase_id=doc.id,product_name='Tablet',line_no=1,status='POSTED',
                           quantity=2,quantity_free=0,raw={'quantity':'2.5','free':'0.5','pack':'10x1x15'},
                           receipt_decision={'resolved':True,'received_base_units':45,'base_unit':'TABLET','master_pack_equivalent':'3'},
                           line_total=100,taxable_value=100,gst_amount=5,cgst_amount=Decimal('2.5'),sgst_amount=Decimal('2.5'),
                           igst_amount=0,landed_total=105,gst_rate=5,batch_no='BATCH1',expiry_date=date(2028,3,1)))
        db.add(PurchaseItem(purchase_id=doc.id,product_name='Not received',line_no=2,status='NEEDS_REVIEW',quantity=1,line_total=50))
    db.commit()
    params=RANGE|{'supplier':str(supplier.id),'invoice_ids':str(docs[0].id)}
    report=reports.generate(db,'supplier-purchases',params)
    assert len(report['rows'])==2
    received,pending=report['rows']
    assert received['paid']==Decimal('2.5') and received['free']==Decimal('0.5')
    assert received['received']==45 and received['equivalent']==3
    assert pending['received'] is None and pending['landed'] is None
    assert report['totals']['landed']==105 and report['totals']['gst']==5
    assert received['difference']==-1 and received['gstin']==''
    from app.services import report_document
    from openpyxl import load_workbook
    import pymupdf
    text=report_document.text_document(report)
    assert 'Billed quantity: 2.5' in text and 'Free quantity: 0.5' in text
    pdf=pymupdf.open(stream=report_document.pdf_bytes(report),filetype='pdf')
    assert 'TOTALS (received lines only)' in ''.join(page.get_text() for page in pdf)
    workbook=load_workbook(BytesIO(report_document.excel_bytes(report)))
    index=next(i for i,c in enumerate(report['columns'],1) if c['key']=='paid')
    assert workbook.active.cell(5,index).value==2.5
    assert workbook.active.cell(5,index).number_format=='#,##0.######'
    with pytest.raises(reports.ReportError,match='do not belong'):
        reports.generate(db,'supplier-purchases',params|{'invoice_ids':str(docs[2].id)})
    with pytest.raises(reports.ReportError,match='valid invoices'):
        reports.generate(db,'supplier-purchases',params|{'invoice_ids':'1,wrong'})
    login(client)
    result=client.get('/reports/api/purchase-invoices',params={'supplier':supplier.id,'period':'custom','from_date':RANGE['from'],'to_date':RANGE['to']})
    assert result.status_code==200
    assert [d['id'] for d in result.json()['invoices']]==[docs[0].id,docs[1].id]
    old=Purchase(supplier_id=supplier.id,invoice_no='OLD-DRAFT',status='DRAFT',purchase_date=datetime(2010,1,1),total=10)
    db.add(old);db.flush()
    db.add(PurchaseItem(purchase_id=old.id,product_name='Old draft item',status='NEEDS_REVIEW',quantity=1,line_total=10))
    db.commit()
    all_invoices=client.get('/reports/api/purchase-invoices',params={'supplier':supplier.id})
    assert all_invoices.status_code==200
    assert {d['id'] for d in all_invoices.json()['invoices']}=={old.id,docs[0].id,docs[1].id}
    audit=reports.generate(db,'supplier-purchases',{'supplier':str(supplier.id),'invoice_ids':str(old.id)})
    assert len(audit['rows'])==1 and audit['rows'][0]['invoice']=='OLD-DRAFT'
    assert audit['rows'][0]['received'] is None and audit['totals']['landed']==0


def test_presets_are_calendar_periods():
    today=date(2026,9,30)
    assert reports.preset_dates('this_week',today)==(date(2026,9,28),today)
    assert reports.preset_dates('last_week',today)==(date(2026,9,21),date(2026,9,27))
    assert reports.preset_dates('last_month',date(2026,3,1))==(date(2026,2,1),date(2026,2,28))


def test_summary_dynamic_categories_reconciles_invoice_and_bill_count(db):
    a,ab=stocked(db,'Medicine','Medicine');b,bb=stocked(db,'Baby','Baby Care')
    sale=sales_service.create_sale(db,lines=[{'item_id':a.id,'batch_id':ab.id,'quantity':2},{'item_id':b.id,'batch_id':bb.id,'quantity':1}],discount='2.50',voucher='1')
    sale.sale_date=datetime(2026,9,1,20);db.commit()
    cancelled=sell(db,a,ab);cancelled.payment_status='CANCELLED';db.commit()
    report=reports.generate(db,'sales-summary',RANGE)
    assert {'cat:MEDICINE','cat:BABY_CARE'}<={c['key'] for c in report['columns']}
    assert 'cat:PHARMA' not in {c['key'] for c in report['columns']}
    assert sum((r['value'] for r in report['rows']),Decimal(0))==sale.total
    assert sum((r['bills'] for r in report['rows']))==1
    for row in report['rows']:
        assert row['cat:MEDICINE']+row['cat:BABY_CARE']==row['value']
    assert report['footer'][0]['value']==sale.total
    # UTC 20:00 belongs to Sep 2 in the store timezone.
    assert report['rows'][0]['date']=='2026-09-02'


def test_date_boundaries_and_void_register(db):
    item,batch=stocked(db)
    first=sell(db,item,batch,when=datetime(2026,8,31,18,30))
    sell(db,item,batch,when=datetime(2026,9,1,18,30))
    void=sell(db,item,batch,when=datetime(2026,9,1,10));void.payment_status='CANCELLED';db.commit()
    r=reports.generate(db,'bill-register',{'period':'custom','from':'2026-09-01','to':'2026-09-01'})
    assert [row['invoice'] for row in r['rows']]==[first.invoice_no]
    r=reports.generate(db,'void-bills',RANGE)
    assert [row['invoice'] for row in r['rows']]==[void.invoice_no]


def test_item_filters_and_only_requested_columns_are_returned(db):
    a,ab=stocked(db,'Match','Medicine');b,bb=stocked(db,'Other','Cosmetics');sell(db,a,ab,qty=2);sell(db,b,bb)
    result=reports.generate(db,'item-sales',RANGE|{'category':'Medicine','manufacturer':'Test Brand'},['item','quantity','value'])
    assert len(result['rows'])==1 and result['rows'][0]['quantity']==2
    assert set(result['rows'][0])=={'item','quantity','value'}
    assert result['totals']['value']==20


@pytest.mark.parametrize('rid',[r['id'] for r in reports.CATALOG])
def test_every_report_queries_existing_database(client,db,rid):
    item,batch=stocked(db);sell(db,item,batch)
    supplier=Supplier(name='Report Supplier');db.add(supplier);db.flush()
    purchase=Purchase(supplier_id=supplier.id,invoice_no='PUR-REPORT',total=50,status='POSTED',purchase_date=datetime(2026,9,1,12));db.add(purchase);db.flush()
    db.add(PurchaseItem(purchase_id=purchase.id,item_id=item.id,product_name=item.name,quantity=10,rate=5,line_total=50))
    db.commit();login(client)
    params={'as_of':date.today().isoformat()} if 'as_of' in reports.BY_ID[rid]['filters'] else dict(RANGE)
    if rid=='customer-history':   # one customer's bills: the customer is required
        assert client.post('/reports/api/generate',json={'report':rid,'parameters':params}).status_code==400
        customer=Customer(name='Report Customer',mobile='9000000001',customer_id='CUST-R1');db.add(customer);db.commit()
        params['customer_id']=customer.id
    if rid=='rack-history':       # one rack's days: the rack is required
        assert client.post('/reports/api/generate',json={'report':rid,'parameters':params}).status_code==400
        from app.services import location_service
        params['rack']=str(location_service.create_rack(db,code='R-T1').id);db.commit()
    response=client.post('/reports/api/generate',json={'report':rid,'parameters':params})
    assert response.status_code==200,response.text
    assert response.json()['id']==rid and 'token' in response.json()


def test_historical_stock_uses_ledger_instead_of_current_projection(db):
    item,batch=stocked(db)
    movement=db.scalar(select(InventoryMovement).where(InventoryMovement.batch_id==batch.id))
    movement.created_at=datetime(2026,8,1,12)
    from app.services.stock_ledger import post
    post(db,batch,'ADJUSTMENT_OUT',3,reason='Later')
    later=db.scalar(select(InventoryMovement).where(InventoryMovement.batch_id==batch.id).order_by(InventoryMovement.id.desc()))
    later.created_at=datetime(2026,9,15,12);batch.created_at=datetime(2026,8,1);db.commit()
    old=reports.generate(db,'batch-stock',{'as_of':'2026-09-01'})
    assert old['rows'][0]['quantity']==20 and batch.quantity==17
    recent=reports.generate(db,'batch-stock',{'as_of':'2026-09-30'})
    assert recent['rows'][0]['quantity']==17


def test_exports_are_selected_snapshot_and_real_files(client,db):
    item,batch=stocked(db,name='=Unsafe Name');sale=sell(db,item,batch)
    login(client)
    r=client.post('/reports/api/generate',json={'report':'item-sales','parameters':RANGE,'columns':['item','quantity','value']}).json()
    # Change SQL after generation: exports must still match the generated document.
    sell(db,item,batch,qty=2)
    base='/reports/api/document/'+r['token']
    csv=client.get(base+'.csv');assert csv.status_code==200
    assert "'=Unsafe Name" in csv.text and 'Customer' not in csv.text
    from openpyxl import load_workbook
    wb=load_workbook(BytesIO(client.get(base+'.xlsx').content));ws=wb.active
    assert [c.value for c in ws[4]]==['Item','Quantity','Sales Value']
    assert ws['B5'].value==1 and ws['C5'].value==10
    pdf=client.get(base+'.pdf');assert pdf.status_code==200 and pdf.content.startswith(b'%PDF')


def test_profit_permission_is_enforced_on_catalog_and_generation(client,db):
    login(client)
    user=db.scalar(select(User).order_by(User.id))
    role=Role(name='Report-only');permission=db.scalar(select(Permission).where(Permission.code=='reports.sales'));role.permissions=[permission];db.add(role);db.flush();user.role_id=role.id;db.commit()
    catalog=client.get('/reports/api/catalog').json()['reports']
    assert 'profit' not in [r['id'] for r in catalog]
    assert client.post('/reports/api/generate',json={'report':'profit','parameters':RANGE}).status_code==403


@pytest.mark.parametrize('params',[{'period':'garbage'},{'period':'custom','from':'2026-09-10','to':'2026-09-01'}, {'period':'custom','from':'invalid','to':'2026-09-01'},{'category':'unsupported'}])
def test_invalid_parameters_refused(client,params):
    login(client)
    response=client.post('/reports/api/generate',json={'report':'sales-summary','parameters':params})
    assert response.status_code==400,response.text


def test_refunds_profit_payments_and_stock_losses(db):
    from app.services import refund_service
    item,batch=stocked(db)
    sale=sell(db,item,batch,qty=4,when=datetime(2026,9,2,12),payment_mode='SPLIT',payments=[{'mode':'CASH','amount':'20'},{'mode':'UPI','amount':'20'}])
    ret=refund_service.create_return(db,sale,lines=[{'sale_item_id':sale.items[0].id,'quantity':1,'disposition':'RESTOCK'}],refund_method='CASH',reason_code='WRONG_ITEM')
    ret.processed_at=datetime(2026,9,3,12);db.commit()
    damaged=refund_service.create_return(db,sale,lines=[{'sale_item_id':sale.items[0].id,'quantity':1,'disposition':'DAMAGE'}],refund_method='UPI',reason_code='DAMAGED')
    damaged.processed_at=datetime(2026,9,4,12);db.commit()
    profit=reports.generate(db,'profit',RANGE)
    assert profit['totals']['revenue']==20
    assert profit['totals']['cogs']==10  # All returns reverse sale COGS; damaged returns are a separate loss.
    assert profit['totals']['profit']==10 and profit['totals']['margin']==50
    detailed=reports.generate(db,'profit',RANGE,['stock_loss','net_profit'])
    assert detailed['totals']['stock_loss']==5 and detailed['totals']['net_profit']==5
    payments=reports.generate(db,'payments',RANGE)
    assert payments['totals']['collected']==40 and payments['totals']['refunded']==20
    assert payments['totals']['value']==20
    returns=reports.generate(db,'sales-returns',RANGE)
    assert returns['totals']['quantity']==2 and returns['totals']['value']==20
    losses=reports.generate(db,'stock-loss',RANGE)
    assert any(r['type']=='DAMAGE' and r['quantity']==-1 for r in losses['rows'])


def test_purchase_filters_returns_and_discount_allocation(db):
    from app.models import PurchaseReturn
    a,ab=stocked(db,'Purchased Medicine','Medicine');b,bb=stocked(db,'Purchased Cosmetic','Cosmetics')
    supplier=Supplier(name='Supplier One');other=Supplier(name='Supplier Two');db.add_all([supplier,other]);db.flush()
    doc=Purchase(supplier_id=supplier.id,invoice_no='INV-1',total=27,status='POSTED',purchase_date=datetime(2026,9,1,12));db.add(doc);db.flush()
    db.add_all([PurchaseItem(purchase_id=doc.id,item_id=a.id,product_name=a.name,quantity=2,rate=10,line_total=20),PurchaseItem(purchase_id=doc.id,item_id=b.id,product_name=b.name,quantity=1,rate=10,line_total=10)])
    ret=PurchaseReturn(purchase_id=doc.id,supplier_id=supplier.id,item_id=a.id,batch_id=ab.id,product_name=a.name,quantity=1,value=9,return_date=datetime(2026,9,5,12),reason='Damaged')
    db.add(ret);db.commit()
    report=reports.generate(db,'purchase-summary',RANGE|{'supplier':str(supplier.id),'category':'Medicine'})
    assert report['totals']['value']==18 and report['totals']['bills']==1
    assert reports.generate(db,'purchase-register',RANGE|{'supplier':str(other.id)})['rows']==[]
    returned=reports.generate(db,'purchase-returns',RANGE|{'supplier':str(supplier.id),'invoice':'INV-1'})
    assert returned['rows'][0]['value']==9 and returned['rows'][0]['item']==a.name


def test_low_stock_includes_products_without_any_batch(db):
    item=inv.create_item(db,name='Never received',category='Medicine',reorder_level=10);db.commit()
    r=reports.generate(db,'low-stock',{'as_of':reports.options(db)['today']})
    assert r['rows'][0]['item']==item.name and r['rows'][0]['quantity']==0
    assert r['rows'][0]['shortfall']==10


def test_sales_summary_detail_has_only_bill_document_columns(db):
    item,batch=stocked(db);sale=sell(db,item,batch)
    result=reports.generate(db,'sales-summary',RANGE|{'view':'detail'},['date','invoice','value'])
    assert result['rows'][0]['invoice']==sale.invoice_no
    assert not any(c['key'].startswith('cat:') for c in result['columns'])


def test_export_snapshot_is_scoped_to_user(client,db):
    login(client)
    r=client.post('/reports/api/generate',json={'report':'bill-register','parameters':RANGE}).json()
    from app.security import hash_password
    from fastapi.testclient import TestClient
    from app.main import app
    admin=db.scalar(select(User).order_by(User.id))
    user=User(employee_id='REPORT-USER',username='report-user',full_name='Report User',password_hash=hash_password('report-test-password'),role_id=admin.role_id)
    db.add(user);db.commit()
    with TestClient(app) as second:
        login(second,username=user.username,password='report-test-password')
        assert second.get('/reports/api/document/'+r['token']+'.csv').status_code==404



def test_void_report_reads_recorded_reason_from_audit(db):
    item,batch=stocked(db);sale=sell(db,item,batch)
    sales_service.void_sale(db,sale,reason='Wrong invoice');db.commit()
    report=reports.generate(db,'void-bills',RANGE)
    assert report['rows'][0]['reason']=='Wrong invoice'


def test_customer_filter_uses_the_picked_customer_and_lookup_is_live(client,db):
    item,batch=stocked(db)
    ali=Customer(name='Ali Khan',mobile='9111111111',customer_id='C-ALI');alia=Customer(name='Alia Begum',mobile='9222222222',customer_id='C-ALIA')
    db.add_all([ali,alia]);db.commit()
    for c in (ali,alia):
        from app.services import sales_service
        sales_service.create_sale(db,lines=[{'item_id':item.id,'quantity':1}],customer_id=c.id)
    db.commit();login(client)
    found=client.get('/reports/api/lookup',params={'kind':'customer','q':'ali'}).json()['results']
    assert {r['label'] for r in found}=={'Ali Khan','Alia Begum'}
    assert client.get('/reports/api/lookup',params={'kind':'item','q':item.name[:4]}).json()['results'][0]['id']==item.id
    typed=reports.generate(db,'bill-register',{'period':'today','customer':'Ali'})
    picked=reports.generate(db,'bill-register',{'period':'today','customer':'Ali Khan','customer_id':str(ali.id)})
    assert len(typed['rows'])==2 and [r['customer'] for r in picked['rows']]==['Ali Khan']
    by_mobile=reports.generate(db,'customer-returns',{'period':'today','customer':'92222'})
    assert by_mobile['rows']==[]
    history=reports.generate(db,'customer-history',{'period':'today','customer':'Ali Khan','customer_id':str(ali.id)})
    assert len(history['rows'])==1
