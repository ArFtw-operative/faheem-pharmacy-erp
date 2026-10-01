"""Sales / POS / billing pages and JSON API."""
from __future__ import annotations

import csv
import io
from datetime import datetime, timedelta

from fastapi import APIRouter, Depends, HTTPException, Query, Request
from fastapi.responses import RedirectResponse, Response
from sqlalchemy import select
from sqlalchemy.orm import Session

from app import audit
from app.routing import OffloadRoute
from app.database import get_db
from app.deps import client_ip, require_login, require_permission
from app.models import (
    RETURN_DISPOSITIONS,
    RETURN_REASONS,
    Customer,
    Sale,
    User,
)
from app.services import (
    business_time,
    customer_service,
    invoice_kit,
    invoice_render,
    parking_service,
    refund_service,
    sales_service,
)
from app.services import settings_service, units
from app.services.printing import get_backend
from app.web import render

router = APIRouter(tags=["sales"], route_class=OffloadRoute)


def _history_filters(db: Session, params: dict) -> dict:
    """Sales-history filters from query params, with periods in the pharmacy's timezone."""
    from datetime import date as _date

    from app.services import reports_service

    tz = business_time.timezone_name(db)
    period = params.get("period") or ""
    start = end = None
    if period in ("today", "week", "month", "year"):
        start, end = reports_service.period_range(period, tz)
    elif period == "custom":
        def local(day: str, add: int = 0):
            try:
                d = _date.fromisoformat(day)
            except (TypeError, ValueError):
                return None
            from zoneinfo import ZoneInfo

            moment = datetime(d.year, d.month, d.day, tzinfo=ZoneInfo(tz)) + timedelta(days=add)
            return moment.astimezone(ZoneInfo("UTC")).replace(tzinfo=None)

        start = local(params.get("date_from") or "")
        end = local(params.get("date_to") or "", add=1)
    return {
        "q": params.get("q") or "",
        "start": start,
        "end": end,
        "payment": params.get("payment") or "",
        "status": params.get("status") or "",
        "customer_type": params.get("ctype") or "",
    }


def _visible_sales_scope(user: User) -> int | None:
    """Return user_id to restrict to, or None for unrestricted."""
    from app.permissions import has_permission

    if has_permission(user, "sales.view_history"):
        return None
    return user.id


@router.post("/api/pos/park")
async def api_park_sale(
    request: Request,
    db: Session = Depends(get_db),
    user: User = Depends(require_permission("sales.create")),
):
    payload = await request.json()
    data = payload.get("payload") or {}
    if not data or not data.get("cart"):
        raise HTTPException(400, "There is nothing to park")
    customer_id = payload.get("customer_id") or (data.get("customer") or {}).get("id")
    if customer_id and db.get(Customer, int(customer_id)) is None:
        customer_id = None
    try:
        parked = parking_service.park_sale(
            db,
            payload=data,
            customer_id=customer_id,
            reason_code=str(payload.get("reason_code") or ""),
            note=str(payload.get("note") or ""),
            user=user,
            ip_address=client_ip(request),
        )
        db.commit()
    except parking_service.ParkingError as exc:
        db.rollback()
        raise HTTPException(400, str(exc))
    return {
        "ok": True,
        "id": parked.id,
        "park_reference": parked.park_reference,
        "total": parking_service.parked_payload(db, parked)["total"],
        "items": parking_service.parked_payload(db, parked)["items"],
    }


@router.get("/api/pos/parked")
def api_parked_list(
    request: Request,
    q: str = "",
    scope: str = "all",
    db: Session = Depends(get_db),
    user: User = Depends(require_permission("sales.create")),
):
    user_id = user.id if scope == "mine" else None
    business_date = business_time.current_business_date(db) if scope == "today" else None
    rows = parking_service.list_parked(
        db, q=q, status=("PARKED", "CLAIMED"), user_id=user_id, business_date=business_date
    )
    return {
        "count": parking_service.count_parked(db),
        "parked": [parking_service.parked_payload(db, p) for p in rows],
    }


@router.get("/api/pos/parked/count")
def api_parked_count(
    db: Session = Depends(get_db),
    user: User = Depends(require_permission("sales.create")),
):
    return {"count": parking_service.count_parked(db)}


@router.post("/api/pos/parked/{parked_id}/resume")
def api_resume_parked(
    parked_id: int,
    request: Request,
    db: Session = Depends(get_db),
    user: User = Depends(require_permission("sales.create")),
):
    parked = parking_service.get_parked(db, parked_id)
    if parked is None:
        raise HTTPException(404, "Parked sale not found")
    try:
        parking_service.resume_parked(db, parked, user=user, ip_address=client_ip(request))
        db.commit()
    except parking_service.ParkedClaimConflict as exc:
        db.rollback()
        raise HTTPException(
            409,
            detail={
                "message": "This parked sale was already resumed elsewhere",
                "parked": parking_service.parked_payload(db, exc.parked) if exc.parked else None,
            },
        )
    return {"ok": True, "parked": parking_service.parked_payload(db, parked)}


@router.post("/api/pos/parked/{parked_id}/release")
def api_release_parked(
    parked_id: int,
    request: Request,
    db: Session = Depends(get_db),
    user: User = Depends(require_permission("sales.create")),
):
    parked = parking_service.get_parked(db, parked_id)
    if parked is None:
        raise HTTPException(404, "Parked sale not found")
    try:
        parking_service.release_parked(db, parked, user=user, ip_address=client_ip(request))
        db.commit()
    except parking_service.ParkingError as exc:
        db.rollback()
        raise HTTPException(400, str(exc))
    return {"ok": True, "parked": parking_service.parked_payload(db, parked)}


@router.post("/api/pos/parked/{parked_id}/discard")
async def api_discard_parked(
    parked_id: int,
    request: Request,
    db: Session = Depends(get_db),
    user: User = Depends(require_permission("sales.create")),
):
    parked = parking_service.get_parked(db, parked_id)
    if parked is None:
        raise HTTPException(404, "Parked sale not found")
    payload = await request.json()
    try:
        parking_service.discard_parked(
            db, parked,
            reason=str(payload.get("reason_code") or ""),
            note=str(payload.get("note") or ""),
            user=user, ip_address=client_ip(request),
        )
        db.commit()
    except parking_service.ParkingError as exc:
        db.rollback()
        raise HTTPException(400, str(exc))
    return {"ok": True}


@router.get("/api/customers/search")
def api_customer_search(
    q: str = "",
    db: Session = Depends(get_db),
    user: User = Depends(require_permission("sales.create")),
):
    if not q.strip():
        return {"customers": []}
    customers, _ = customer_service.search_customers(db, q=q, limit=10)
    return {"customers": [customer_service.customer_payload(c) for c in customers]}


@router.get("/api/customers/lookup")
def api_customer_lookup(
    mobile: str = "",
    db: Session = Depends(get_db),
    user: User = Depends(require_permission("sales.create")),
):
    """Existing customer for a mobile number (formatting and +91 ignored)."""
    customer = customer_service.find_by_mobile(db, mobile) if len(customer_service.mobile_key(mobile)) >= 10 else None
    return {"customer": customer_service.customer_payload(customer) if customer else None}


@router.get("/api/customers/mobile-search")
def api_customer_mobile_search(
    q: str = "",
    db: Session = Depends(get_db),
    user: User = Depends(require_permission("sales.create")),
):
    """Live POS search: customers whose mobile contains the digits typed so far."""
    return {"customers": [customer_service.customer_payload(c) for c in customer_service.search_by_mobile(db, q)]}


def _resolve_bill_customer(db: Session, payload: dict, user: User, ip_address: str) -> int | None:
    """Customer for a POS bill: the chosen record, or the name/mobile typed at the counter.

    A typed mobile that already exists links that customer (never a duplicate);
    otherwise a new customer is created in the same transaction as the sale.
    """
    from app.permissions import has_permission

    if payload.get("customer_id"):
        return payload["customer_id"]
    typed = payload.get("new_customer") or {}
    name = " ".join(str(typed.get("name") or "").split())[:150]
    mobile = str(typed.get("mobile") or "").strip()[:20]
    if not name and not mobile:
        return None
    if mobile:
        if len(customer_service.mobile_key(mobile)) != 10:
            raise HTTPException(400, "Enter a 10-digit mobile number for the customer")
        existing = customer_service.find_by_mobile(db, mobile)
        if existing is not None:
            return existing.id
    if not name:
        raise HTTPException(400, "Enter the customer's name to save them with this bill")
    if not has_permission(user, "customers.create"):
        raise HTTPException(403, "You are not allowed to add customers")
    customer = customer_service.create_customer(
        db, name=name, mobile=mobile, user=user, ip_address=ip_address,
        customer_type=payload.get("customer_type") or "WALK_IN",
    )
    return customer.id


@router.post("/api/customers")
async def api_customer_create(
    request: Request,
    db: Session = Depends(get_db),
    user: User = Depends(require_permission("customers.create")),
):
    payload = await request.json()
    mobile = str(payload.get("mobile") or "").strip()
    name = str(payload.get("name") or "").strip()
    if not name:
        raise HTTPException(400, "Customer name is required")
    if not mobile:
        raise HTTPException(400, "Mobile number is required")

    existing = customer_service.find_by_mobile(db, mobile)
    if existing is not None:
        raise HTTPException(
            409,
            detail={
                "message": f"A customer with mobile {mobile} already exists.",
                "customer": customer_service.customer_payload(existing),
            },
        )

    try:
        customer = customer_service.create_customer(
            db,
            name=name,
            user=user,
            ip_address=client_ip(request),
            mobile=mobile,
            alternate_mobile=str(payload.get("alternate_mobile") or ""),
            email=str(payload.get("email") or ""),
            gender=str(payload.get("gender") or ""),
            date_of_birth=payload.get("date_of_birth") or None,
            doctor_name=str(payload.get("doctor_name") or ""),
            address=str(payload.get("address") or ""),
            city=str(payload.get("city") or ""),
            state=str(payload.get("state") or ""),
            pincode=str(payload.get("pincode") or ""),
            notes=str(payload.get("notes") or ""),
            customer_type=payload.get("customer_type") or payload.get("type") or "WALK_IN",
        )
        db.commit()
    except customer_service.CustomerError as exc:
        db.rollback()
        raise HTTPException(400, str(exc))
    return {"ok": True, "customer": customer_service.customer_payload(customer)}


@router.post("/api/customers/{customer_id}/type")
async def api_customer_set_type(
    customer_id: int,
    request: Request,
    db: Session = Depends(get_db),
    user: User = Depends(require_permission("sales.create")),
):
    """In-place category change from the POS customer strip."""
    customer = db.get(Customer, customer_id)
    if customer is None or not customer.is_active:
        raise HTTPException(404, "Customer not found")
    payload = await request.json()
    customer_service.update_customer(
        db,
        customer,
        user=user,
        ip_address=client_ip(request),
        customer_type=payload.get("customer_type"),
    )
    db.commit()
    return {"ok": True, "customer": customer_service.customer_payload(customer)}


@router.get("/api/sales/{sale_id}/invoice-view")
def api_sale_invoice_view(
    sale_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(require_login),
):
    """The saved invoice mapped to the customer-invoice kit's data contract."""
    from app.permissions import has_permission

    if not (has_permission(user, "sales.create") or has_permission(user, "sales.view")):
        raise HTTPException(403, "Not allowed to view invoices")
    sale = sales_service.get_sale(db, sale_id)
    if sale is None:
        raise HTTPException(404, "Sale not found")
    if sale.payment_status == "CANCELLED":
        raise HTTPException(400, "This invoice was cancelled")
    return invoice_kit.build_invoice_view(db, sale)


@router.get("/api/sales/recent")
def api_recent_sales(
    limit: int = Query(8, ge=1, le=50),
    db: Session = Depends(get_db),
    user: User = Depends(require_permission("sales.view_own")),
):
    scope = _visible_sales_scope(user)
    sales, _ = sales_service.search_sales(db, user_id=scope, limit=limit)
    return {
        "sales": [
            {
                "id": s.id,
                "invoice_no": s.invoice_no,
                "total": str(s.total),
                "payment_mode": s.payment_mode,
                "payment_label": sales_service.payment_label(s),
                "payment_status": s.payment_status,
                "date": s.sale_date.strftime("%d %b %H:%M"),
                "customer": s.customer.name if s.customer else "Walk-in",
                "customer_type": s.customer_type or "WALK_IN",
                "items": len(s.items),
            }
            for s in sales
        ]
    }


@router.post("/api/sales")
async def api_create_sale(
    request: Request,
    db: Session = Depends(get_db),
    user: User = Depends(require_permission("sales.create")),
):
    from sqlalchemy.exc import IntegrityError

    from app.permissions import has_permission

    payload = await request.json()
    lines = payload.get("lines") or []
    if not has_permission(user, "sales.discount"):
        payload["discount"] = 0
        payload["discount_pct"] = None
        payload["voucher"] = 0
        for line in lines:
            line.pop("discount", None)
            line.pop("discount_pct", None)
    request_id = payload.get("client_request_id") or None
    parked_id = payload.get("parked_id")

    try:
        customer_id = _resolve_bill_customer(db, payload, user, client_ip(request))
        sale = sales_service.create_sale(
            db,
            lines=lines,
            user=user,
            customer_id=customer_id,
            customer_type=payload.get("customer_type", "WALK_IN"),
            payment_mode=payload.get("payment_mode", "CASH"),
            payments=payload.get("payments") or None,
            discount=payload.get("discount", 0),
            discount_pct=payload.get("discount_pct"),
            voucher=payload.get("voucher", 0),
            notes=payload.get("notes", ""),
            round_off_mode=settings_service.get_setting(db, "round_off_mode", "NEAREST_RUPEE"),
            ip_address=client_ip(request),
            cash_received=payload.get("cash_received"),
            client_request_id=request_id,
            invoice_type=payload.get("invoice_type") or "INVENTORY",
        )
        db.commit()
    except IntegrityError:
        # Concurrent duplicate submission: return the sale that won the race.
        db.rollback()
        if not request_id:
            raise HTTPException(400, "Duplicate sale submission")
        sale = db.scalar(select(Sale).where(Sale.client_request_id == request_id))
        if sale is None:
            raise HTTPException(400, "Duplicate sale submission")
    except (sales_service.SaleError, customer_service.CustomerError) as exc:
        db.rollback()
        raise HTTPException(400, str(exc))
    except HTTPException:
        db.rollback()
        raise

    if parked_id:
        parked = parking_service.get_parked(db, int(parked_id))
        if parked is not None and parked.status in ("PARKED", "CLAIMED"):
            parking_service.complete_parked(db, parked, sale.id, user=user)
            db.commit()

    return {
        "ok": True,
        "sale_id": sale.id,
        "invoice_no": sale.invoice_no,
        "total": str(sale.total),
        "tendered": str(sale.tendered_amount) if sale.tendered_amount is not None else None,
        "change": str(sale.change_amount) if sale.change_amount is not None else None,
        "customer": customer_service.customer_payload(sale.customer) if sale.customer else None,
    }


@router.get("/sales")
def sales_list(user: User = Depends(require_permission("sales.view_own"))):
    """The bill register lives in the ERP Sales module now."""
    return RedirectResponse("/app/sales", status_code=303)


def _visible_sale(db: Session, sale_id: int, user: User) -> Sale:
    sale = sales_service.get_sale(db, sale_id)
    if sale is None:
        raise HTTPException(404, "Sale not found")
    scope = _visible_sales_scope(user)
    if scope is not None and sale.user_id != scope:
        raise HTTPException(403, "You can only view your own sales")
    return sale


INVOICE_SIZES = ("A4", "A5", "LETTER", "THERMAL80", "THERMAL58", "ERP")


@router.get("/sales/{sale_id}/invoice")
def sale_invoice_document(
    sale_id: int,
    request: Request,
    size: str = "A4",
    mono: int = 0,
    expiry: int = 1,
    autoprint: int = 0,
    db: Session = Depends(get_db),
    user: User = Depends(require_permission("sales.view_own")),
):
    """The invoice rendered by the studio engine — shown inside the ERP tab (POS or Sales)."""
    sale = _visible_sale(db, sale_id, user)
    size = size.upper() if size.upper() in INVOICE_SIZES else "A4"
    if size == "ERP":                     # the plain ERP document of this one bill
        from app.services import report_document

        return render(request, "invoice_text.html", db, {"sale": sale, "text": report_document.sale_text(db, sale),
                                                          "autoprint": bool(autoprint)}, user=user)
    return render(request, "invoice_print.html", db, {
        "sale": sale, "invoice_data": invoice_kit.build_invoice_view(db, sale),
        "options": {"size": size, "mono": bool(mono) or size.startswith("THERMAL"), "expiry": bool(expiry), "autoprint": bool(autoprint)},
    }, user=user)


@router.get("/api/erp/sales/{sale_id}/export.{fmt}")
def sale_invoice_export(sale_id: int, fmt: str, db: Session = Depends(get_db),
                        user: User = Depends(require_permission("sales.view_own"))):
    """One invoice as Excel or CSV: header facts, lines, totals and payments."""
    sale = _visible_sale(db, sale_id, user)
    if fmt not in ("csv", "xlsx"):
        raise HTTPException(400, "Choose Excel or CSV")
    head = [("Invoice", sale.invoice_no), ("Type", "Manual bill" if sale.invoice_type == "MANUAL" else "Stock bill"),
            ("Date", sale.sale_date.strftime("%d-%m-%Y %H:%M")), ("Customer", sale.customer.name if sale.customer else "Walk-in"),
            ("Mobile", sale.customer.mobile if sale.customer else ""), ("Status", "Voided" if sale.payment_status == "CANCELLED" else "Completed")]
    cols = ["Sr", "Item", "Pack", "Batch", "Expiry", "Qty", "Unit", "Rate", "Discount", "Amount"]
    lines = [[i + 1, l.product_name, l.pack_size, l.batch_no, l.expiry_date.strftime("%m/%Y") if l.expiry_date else "",
              l.quantity, (l.base_unit or "UNIT").lower(), float(l.rate), float(l.discount), float(l.line_total)]
             for i, l in enumerate(sorted(sale.items, key=lambda l: (l.line_no or 0, l.id)))]
    totals = [("Items total", float(sale.subtotal)), ("Bill discount", float(sale.discount or 0)),
              ("Round off", float(sale.round_off or 0)), ("Bill amount", float(sale.total))]
    pays = [(f"Paid by {p.mode.title()}", float(p.amount)) for p in sale.payments]
    if sale.tendered_amount is not None:
        pays += [("Received", float(sale.tendered_amount)), ("Change", float(sale.change_amount or 0))]
    audit.record(db, action=audit.A_EXPORT, entity_type="sale", entity_id=sale.invoice_no, user=user,
                 details=f"Invoice exported as {fmt.upper()}", commit=True)
    if fmt == "csv":
        buf = io.StringIO()
        w = csv.writer(buf)
        for k, v in head:
            w.writerow([k, v])
        w.writerow([])
        w.writerow(cols)
        w.writerows(lines)
        w.writerow([])
        for k, v in totals + pays:
            w.writerow([k, f"{v:.2f}"])
        return Response(buf.getvalue().encode("utf-8-sig"), media_type="text/csv",
                        headers={"Content-Disposition": f"attachment; filename={sale.invoice_no}.csv"})
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Invoice"
    ws.append([settings_service.get_profile(db).get("name", "") if isinstance(settings_service.get_profile(db), dict) else ""])
    ws["A1"].font = Font(bold=True, size=13)
    for k, v in head:
        ws.append([k, v])
    ws.append([])
    ws.append(cols)
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for row in lines:
        ws.append(row)
    ws.append([])
    for k, v in totals + pays:
        ws.append([k, v])
    for col, width in zip("ABCDEFGHIJ", (6, 34, 10, 12, 9, 7, 8, 10, 10, 12)):
        ws.column_dimensions[col].width = width
    out = io.BytesIO()
    wb.save(out)
    return Response(out.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename={sale.invoice_no}.xlsx"})


@router.get("/sales/{sale_id}/pdf")
def sale_pdf(
    sale_id: int,
    request: Request,
    db: Session = Depends(get_db),
    user: User = Depends(require_permission("sales.view_own")),
):
    sale = sales_service.get_sale(db, sale_id)
    if sale is None:
        raise HTTPException(404, "Sale not found")
    scope = _visible_sales_scope(user)
    if scope is not None and sale.user_id != scope:
        raise HTTPException(403, "Not permitted")
    profile = settings_service.get_profile(db)
    qr = settings_service.qr_png_bytes(db)
    pdf = invoice_render.build_invoice_pdf(sale, profile, qr)
    audit.record(
        db, action=audit.A_EXPORT, entity_type="sale", entity_id=sale.invoice_no, user=user,
        details="Invoice PDF generated", commit=True,
    )
    return Response(
        pdf,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={sale.invoice_no}.pdf"},
    )


@router.post("/sales/{sale_id}/print")
def sale_print(
    sale_id: int,
    request: Request,
    db: Session = Depends(get_db),
    user: User = Depends(require_permission("billing.print")),
):
    sale = sales_service.get_sale(db, sale_id)
    if sale is None:
        raise HTTPException(404, "Sale not found")
    profile = settings_service.get_profile(db)
    backend = get_backend()
    result = backend.print_text(invoice_render.invoice_text(sale, profile))
    audit.record(
        db,
        action=audit.A_PRINT,
        entity_type="sale",
        entity_id=sale.invoice_no,
        user=user,
        details=f"Printed via {result.backend}: {result.message}",
        ip_address=client_ip(request),
        commit=True,
    )
    return {"ok": result.ok if hasattr(result, "ok") else True, "backend": result.backend, "message": result.message}


@router.get("/api/erp/sales/{sale_id}/edit")
def api_sale_edit_payload(sale_id: int, db: Session = Depends(get_db),
                          user: User = Depends(require_permission("sales.void"))):
    """A completed bill reopened in a POS tab: its lines with every sellable batch.

    Availability counts the bill's own units as available again — they go back
    to their batches when the edit is saved."""
    from datetime import date as _date

    from app.models import Batch, Item
    from app.services import inventory_service as inv

    sale = sales_service.get_sale(db, sale_id)
    if sale is None:
        raise HTTPException(404, "Sale not found")
    problem = sales_service.editable_problem(db, sale)
    if problem:
        raise HTTPException(400, problem)
    own: dict[int, int] = {}
    for line in sale.items:
        if line.batch_id:
            own[line.batch_id] = own.get(line.batch_id, 0) + line.quantity
    today = _date.today()
    lines = []
    for line in sorted(sale.items, key=lambda l: (l.line_no or 0, l.id)):
        gross = (line.line_total or 0) + (line.discount or 0)
        pct = float(round(line.discount / gross * 100, 2)) if gross else 0.0
        if not line.item_id:
            lines.append({"manual": True, "name": line.product_name, "qty": line.quantity, "rate": float(line.rate), "disc": pct})
            continue
        item = db.get(Item, line.item_id)
        if item is None:
            raise HTTPException(400, f"{line.product_name} no longer exists in inventory; this bill cannot be edited")
        batches = [b for b in db.scalars(select(Batch).where(Batch.item_id == item.id).order_by(Batch.expiry_date.asc().nulls_last(), Batch.id))
                   if (b.quantity + own.get(b.id, 0)) > 0 and (not units.is_expired(b.expiry_date, today) or b.id in own)]
        pv = inv.packaging_view(item)
        lines.append({
            "item_id": item.id, "code": item.article_id, "name": item.name, "pack_raw": item.pack_size, "upp": item.units_per_pack or 1,
            "loose": bool(item.loose_sale), "base_unit": item.base_unit, "pack_unit": item.pack_unit, "form": item.dosage_form,
            "rack": item.rack or "", "content": pv["content"], "batch_id": line.batch_id, "qty": line.quantity, "disc": pct,
            "batches": [{"id": b.id, "batch_no": b.batch_no, "expiry": b.expiry_date.isoformat() if b.expiry_date else "",
                         "stock": b.quantity + own.get(b.id, 0), "pack_mrp": str(b.mrp), "upp": b.units_per_pack or 1,
                         "unit_mrp": str(units.display_unit_price(b.mrp, b.units_per_pack or 1))} for b in batches],
        })
    items_total = sum((l.line_total or 0 for l in sale.items), 0)
    bill_pct = float(round(sale.discount / items_total * 100, 2)) if items_total and sale.discount else 0.0
    return {
        "sale_id": sale.id, "invoice_no": sale.invoice_no, "invoice_type": sale.invoice_type or "INVENTORY",
        "lines": lines, "customer": customer_service.customer_payload(sale.customer) if sale.customer else None,
        "customer_type": sale.customer_type or "WALK_IN", "discount_pct": bill_pct, "notes": sale.notes or "",
        "payment_mode": sale.payment_mode, "payments": [{"mode": p.mode, "amount": str(p.amount), "reference": p.reference} for p in sale.payments],
        "tendered": str(sale.tendered_amount) if sale.tendered_amount is not None else None, "total": str(sale.total),
    }


@router.put("/api/sales/{sale_id}")
async def api_sale_amend(sale_id: int, request: Request, db: Session = Depends(get_db),
                         user: User = Depends(require_permission("sales.void"))):
    """Save an edited bill: same number, old stock back to its batches, new lines sold, fully audited."""
    from app.permissions import has_permission

    sale = sales_service.get_sale(db, sale_id)
    if sale is None:
        raise HTTPException(404, "Sale not found")
    payload = await request.json()
    lines = payload.get("lines") or []
    if not has_permission(user, "sales.discount"):
        payload["discount"] = 0
        payload["discount_pct"] = None
        for line in lines:
            line.pop("discount", None)
            line.pop("discount_pct", None)
    try:
        customer_id = _resolve_bill_customer(db, payload, user, client_ip(request))
        sales_service.amend_sale(
            db, sale, lines=lines, user=user, customer_id=customer_id,
            customer_type=payload.get("customer_type", sale.customer_type or "WALK_IN"),
            payment_mode=payload.get("payment_mode", "CASH"), payments=payload.get("payments") or None,
            discount=payload.get("discount", 0), discount_pct=payload.get("discount_pct"), notes=payload.get("notes"),
            round_off_mode=settings_service.get_setting(db, "round_off_mode", "NEAREST_RUPEE"),
            cash_received=payload.get("cash_received"), reason=str(payload.get("reason") or "")[:200],
            ip_address=client_ip(request),
        )
        db.commit()
    except (sales_service.SaleError, customer_service.CustomerError) as exc:
        db.rollback()
        raise HTTPException(400, str(exc))
    except HTTPException:
        db.rollback()
        raise
    sale = sales_service.get_sale(db, sale_id)
    return {
        "ok": True, "sale_id": sale.id, "invoice_no": sale.invoice_no, "total": str(sale.total),
        "tendered": str(sale.tendered_amount) if sale.tendered_amount is not None else None,
        "change": str(sale.change_amount) if sale.change_amount is not None else None,
        "customer": customer_service.customer_payload(sale.customer) if sale.customer else None,
    }


@router.get("/api/sales/{sale_id}/refundable")
def api_refundable(
    sale_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(require_permission("sales.refund")),
):
    sale = sales_service.get_sale(db, sale_id)
    if sale is None:
        raise HTTPException(404, "Sale not found")
    return {
        "sale_id": sale.id,
        "invoice_no": sale.invoice_no,
        "customer": sale.customer.name if sale.customer else "Walk-in",
        "total": str(sale.total),
        "payment_mode": sale.payment_mode,
        "refunded_total": str(refund_service.refunded_total(db, sale)),
        "refund_status": refund_service.refund_status(db, sale),
        "lines": refund_service.refundable_lines(db, sale),
        "reason_codes": list(RETURN_REASONS),
        "dispositions": list(RETURN_DISPOSITIONS),
    }


@router.post("/api/sales/{sale_id}/refund")
async def api_create_refund(
    sale_id: int,
    request: Request,
    db: Session = Depends(get_db),
    user: User = Depends(require_permission("sales.refund")),
):
    from app.permissions import has_permission

    sale = sales_service.get_sale(db, sale_id)
    if sale is None:
        raise HTTPException(404, "Sale not found")
    payload = await request.json()
    refund_method = str(payload.get("refund_method") or "CASH").upper()
    try:
        ret = refund_service.create_return(
            db, sale,
            lines=payload.get("lines") or [],
            refund_method=refund_method,
            reason_code=str(payload.get("reason_code") or ""),
            reason_note=str(payload.get("reason_note") or ""),
            disposition=str(payload.get("disposition") or "RESTOCK"),
            manager_threshold=settings_service.get_float(db, "refund_manager_threshold", 0.0),
            allow_override=has_permission(user, "sales.refund_override"),
            user=user,
            ip_address=client_ip(request),
        )
        db.commit()
    except refund_service.RefundError as exc:
        db.rollback()
        raise HTTPException(400, str(exc))
    return {"ok": True, "return": refund_service.return_payload(db, ret)}


@router.get("/api/sales/{sale_id}/refunds")
def api_sale_refunds(
    sale_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(require_permission("sales.view_own")),
):
    sale = sales_service.get_sale(db, sale_id)
    if sale is None:
        raise HTTPException(404, "Sale not found")
    return {
        "refunded_total": str(refund_service.refunded_total(db, sale)),
        "refund_status": refund_service.refund_status(db, sale),
        "returns": [refund_service.return_payload(db, r) for r in refund_service.list_returns(db, sale_id=sale_id)],
    }


@router.get("/sales/returns/{return_id}/pdf")
def return_pdf(
    return_id: int,
    request: Request,
    db: Session = Depends(get_db),
    user: User = Depends(require_permission("sales.refund")),
):
    ret = refund_service.get_return(db, return_id)
    if ret is None:
        raise HTTPException(404, "Refund not found")
    profile = settings_service.get_profile(db)
    pdf = invoice_render.build_refund_pdf(ret, ret.sale, profile)
    audit.record(
        db, action=audit.A_PRINT, entity_type="sale_return", entity_id=ret.return_no,
        user=user, details="Refund receipt generated", ip_address=client_ip(request), commit=True,
    )
    return Response(
        pdf, media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={ret.return_no}.pdf"},
    )


@router.get("/sales/export/{fmt}")
def export_sales(
    fmt: str,
    request: Request,
    db: Session = Depends(get_db),
    user: User = Depends(require_permission("sales.export")),
):
    filters = _history_filters(db, dict(request.query_params))
    sales, _ = sales_service.search_sales(db, user_id=_visible_sales_scope(user), limit=1_000_000, **filters)
    columns = ["invoice_no", "sale_date", "customer", "customer_id", "mobile", "customer_type", "employee", "payment_mode", "payments", "items", "subtotal", "discount", "voucher", "round_off", "total", "status"]
    rows = []
    for s in sales:
        rows.append({
            "invoice_no": s.invoice_no,
            "sale_date": s.sale_date.strftime("%Y-%m-%d %H:%M"),
            "customer": s.customer.name if s.customer else "Walk-in",
            "customer_id": s.customer.customer_id if s.customer else "",
            "mobile": s.customer.mobile if s.customer else "",
            "customer_type": s.customer_type or "WALK_IN",
            "employee": s.user.username if s.user else "",
            "payment_mode": s.payment_mode,
            "payments": sales_service.payment_label(s),
            "items": len(s.items),
            "subtotal": str(s.subtotal),
            "discount": str(s.discount),
            "voucher": str(s.voucher),
            "round_off": str(s.round_off),
            "total": str(s.total),
            "status": "VOID" if s.payment_status == "CANCELLED" else "PAID",
        })
    audit.record(db, action=audit.A_EXPORT, entity_type="sales", entity_id=fmt, user=user, details=f"Exported {len(rows)} sales", commit=True)
    if fmt == "csv":
        buf = io.StringIO()
        writer = csv.DictWriter(buf, fieldnames=columns)
        writer.writeheader()
        writer.writerows(rows)
        return Response(buf.getvalue().encode("utf-8-sig"), media_type="text/csv", headers={"Content-Disposition": "attachment; filename=sales.csv"})
    if fmt == "xlsx":
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "Sales"
        ws.append(columns)
        for row in rows:
            ws.append([row[c] for c in columns])
        out = io.BytesIO()
        wb.save(out)
        return Response(out.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=sales.xlsx"})
    raise HTTPException(400, "Unsupported format")
