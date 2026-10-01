"""Export exactly the selected columns of a generated report snapshot."""
from io import BytesIO, StringIO
import csv
from datetime import timezone
from decimal import Decimal


def safe_text(value):
    text=str(value or '')
    return "'"+text if text.startswith(('=','+','-','@','\t','\r')) else text


def export_rows(report):
    cols=report['columns']
    rows=[[r.get(c['key'],'') for c in cols] for r in report['rows']]
    if report['totals']:
        row=[report['totals'].get(c['key'],'') for c in cols]
        label=next((i for i,c in enumerate(cols) if c['kind']=='text'),None)
        if label is not None: row[label]='TOTAL'
        rows.append(row)
    return cols,rows


def csv_bytes(report):
    cols,rows=export_rows(report);out=StringIO();writer=csv.writer(out)
    writer.writerow([c['label'] for c in cols])
    for row in rows:
        writer.writerow([v if c['kind'] in ('money','number') else safe_text(v) for c,v in zip(cols,row)])
    return out.getvalue().encode('utf-8-sig')


def excel_bytes(report):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    cols,rows=export_rows(report);wb=Workbook();ws=wb.active;ws.title='Report'
    for text in [report['pharmacy'],report['title'].upper(),f"From: {report['from_date']}   To: {report['to_date']}   Generated: {report['generated_at']}"]:
        ws.append([safe_text(text)]);ws.merge_cells(start_row=ws.max_row,start_column=1,end_row=ws.max_row,end_column=len(cols))
    ws.append([c['label'] for c in cols]);ws.freeze_panes='A5';ws.auto_filter.ref=f'A4:{get_column_letter(len(cols))}{4+len(report["rows"])}'
    for row in rows:
        ws.append([float(v) if c['kind'] in ('money','number') and v not in ('',None) else safe_text(v) for c,v in zip(cols,row)])
    line=Side(style='hair',color='D5DCDF')
    for row in ws.iter_rows(min_row=4):
        for cell in row:
            cell.border=Border(bottom=line,right=line)
            if cell.row==4: cell.font=Font(bold=True);cell.fill=PatternFill('solid',fgColor='F4F6F7')
            elif cols[cell.column-1]['kind'] in ('money','number'):
                cell.alignment=Alignment(horizontal='right');cell.number_format='#,##0.00' if cols[cell.column-1]['kind']=='money' else '#,##0'
    if report['totals']:
        for cell in ws[ws.max_row]: cell.font=Font(bold=True);cell.border=Border(top=Side(style='thin',color='1B2327'))
    for i,c in enumerate(cols,1): ws.column_dimensions[get_column_letter(i)].width=28 if c['kind']=='text' else 18
    for f in report['footer']: ws.append([f['label'],float(f['value'])])
    if report['note']: ws.append([report['note']])
    ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A4;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    ws.print_title_rows='1:4';out=BytesIO();wb.save(out);return out.getvalue()


_MONTHS=('Jan','Feb','Mar','Apr','May','Jun','Jul','Aug','Sep','Oct','Nov','Dec')
TEXT_WIDTH_CAP=42          # a very long name is cut, never allowed to break the layout


def _dmy(iso):
    """2026-09-01 → 01-Sep-2026 (anything else unchanged)."""
    text=str(iso or '')
    if len(text)>=10 and text[4]=='-' and text[7]=='-' and text[:4].isdigit():
        return f"{text[8:10]}-{_MONTHS[int(text[5:7])-1]}-{text[:4]}"+text[10:16].replace('T',' ')
    return text


def _cell(c,value):
    if value is None: return '—'
    if value=='': return ''
    if c['kind']=='money':
        from app.utils import format_inr
        try: return format_inr(value,symbol=False).replace('\u2212','-')
        except Exception: return str(value)
    if c['kind']=='number':
        try:
            d=float(value); return f"{int(d):,}" if d==int(d) else f"{d:,.2f}"
        except (TypeError,ValueError): return str(value)
    if c['key'] in ('date','day','from','to','expiry','as_of') or c.get('kind')=='date': return _dmy(value)
    return str(value)


def text_document(report, width_min=78):
    """The report as an ERP text document:

                               FAHEEM PHARMACY
                               ITEM-WISE SALES
    Period: 01-Sep-2026 to 30-Sep-2026
    ----------------------------------------------------------------
    SNo   Item Name        Pack      Qty      MRP    Sales Value
    ----------------------------------------------------------------
    1     Dolo 650         15 TAB     21    32.00         672.00
    ----------------------------------------------------------------
    TOTAL                             21                  672.00
    ----------------------------------------------------------------

    Numbers right-aligned, text left-aligned; cost / profit columns appear
    only when the report's selected columns include them."""
    cols=report['columns']
    rows=[[_cell(c,r.get(c['key'])) for c in cols] for r in report['rows']]
    head=['SNo']+[c['label'] for c in cols]
    body=[[str(i)]+[v if len(v)<=TEXT_WIDTH_CAP else v[:TEXT_WIDTH_CAP-1]+'…' for v in row] for i,row in enumerate(rows,1)]
    total=None
    if report.get('totals'):
        total=['TOTAL']+[_cell(c,report['totals'][c['key']]) if c['key'] in report['totals'] else '' for c in cols]
    kinds=['text']+[c['kind'] for c in cols]
    widths=[max(len(head[i]),*(len(r[i]) for r in body),len(total[i]) if total else 0) for i in range(len(head))]
    widths[0]=max(widths[0],5)
    def line(values):
        parts=[]
        for i,(v,w) in enumerate(zip(values,widths)):
            right=kinds[i] in ('money','number')
            parts.append(v.rjust(w) if right else v.ljust(w))
        return '   '.join(parts).rstrip()
    table_width=max(width_min,sum(widths)+3*(len(widths)-1))
    rule='-'*table_width
    period=(f"As of: {_dmy(report['to_date'])}" if report.get('parameters',{}).get('as_of')
            else f"Period: {_dmy(report['from_date'])} to {_dmy(report['to_date'])}")
    out=[report['pharmacy'].upper().center(table_width).rstrip(),'',report['title'].upper().center(table_width).rstrip(),'',
         period+f"Generated: {report['generated_at']}".rjust(table_width-len(period)),rule,line(head),rule]
    out+=[line(r) for r in body] or ['No records match the selected parameters.']
    out.append(rule)
    if total:
        out+=[line(total),rule]
    for f in report.get('footer') or []:
        out.append(f"{f['label']}: {_cell({'kind':'money','key':''},f['value'])}")
    if report.get('note'):
        import textwrap
        out+=['']+textwrap.wrap(report['note'],table_width)
    return '\n'.join(out)+'\n'


def text_bytes(report):
    return text_document(report).encode('utf-8')


def pdf_bytes(report):
    """PDF of the same text document (monospace), landscape, header repeated on every page."""
    import pymupdf
    lines=text_document(report).split('\n')
    try: rule_at=[i for i,l in enumerate(lines) if l and set(l)=={'-'}]
    except Exception: rule_at=[]
    head=lines[:rule_at[2]+1] if len(rule_at)>=3 else lines[:8]
    rest=lines[len(head):]
    width_chars=max((len(l) for l in lines),default=80)
    doc=pymupdf.open()
    page_w,page_h=842,595
    size=min(9.0,(page_w-48)/(width_chars*0.6))
    step=size*1.35
    per_page=max(10,int((page_h-48)/step)-len(head)-1)
    chunks=[rest[i:i+per_page] for i in range(0,len(rest),per_page)] or [[]]
    for n,chunk in enumerate(chunks,1):
        page=doc.new_page(width=page_w,height=page_h)
        y=24+size
        for l in head+chunk:
            page.insert_text((24,y),l,fontname='cour',fontsize=size)
            y+=step
        page.insert_text((page_w-90,page_h-14),f'Page {n} of {len(chunks)}',fontname='cour',fontsize=7)
    return doc.tobytes()


def sale_text(db, sale, width=80):
    """One bill as a plain ERP document (the individual-invoice view of Sales History)."""
    from zoneinfo import ZoneInfo

    from app.services import business_time, sales_service, settings_service

    tz = ZoneInfo(business_time.timezone_name(db))
    profile = settings_service.get_profile(db)
    name = (profile.get("pharmacy_name") if isinstance(profile, dict) else "") or "FAHEEM PHARMACY"
    when = sale.sale_date.replace(tzinfo=timezone.utc).astimezone(tz)
    title = "SALES INVOICE" + (" (MANUAL BILL)" if (sale.invoice_type or "") == "MANUAL" else "") + (" — VOIDED" if sale.payment_status == "CANCELLED" else "")
    cust = sale.customer
    left = lambda a, b: a + b.rjust(width - len(a))
    cols = [("SNo", 4, "l"), ("Item", 24, "l"), ("Batch", 11, "l"), ("Qty", 8, "r"), ("MRP", 9, "r"), ("Disc.", 8, "r"), ("Amount", 10, "r")]
    def line(vals):
        return " ".join((str(v)[:w].ljust(w) if a == "l" else str(v)[:w].rjust(w)) for (_, w, a), v in zip(cols, vals)).rstrip()
    rule = "-" * width
    out = [name.upper().center(width).rstrip(), title.center(width).rstrip(), "",
           left(f"Invoice: {sale.invoice_no}", f"Date: {when:%d-%b-%Y %I:%M %p}"),
           left(f"Customer: {cust.name if cust else 'Walk-in'}", f"Mobile: {cust.mobile if cust and cust.mobile else '—'}"),
           "", rule, line([c[0] for c in cols]), rule]
    for i, l in enumerate(sorted(sale.items, key=lambda l: (l.line_no or 0, l.id)), 1):
        unit = (l.base_unit or "unit").lower()
        out.append(line([i, l.product_name, l.batch_no or "—", f"{l.quantity} {unit[:3]}", f"{l.mrp:.2f}",
                         f"{l.discount:.2f}" if l.discount else "", f"{l.line_total:.2f}"]))
    out.append(rule)
    items_total = sum((l.line_total for l in sale.items), Decimal(0)) if sale.items else Decimal(0)
    totals = [("Items total", items_total)]
    if sale.discount:
        totals.append(("Bill discount", -sale.discount))
    if sale.round_off:
        totals.append(("Round off", sale.round_off))
    totals.append(("TOTAL", sale.total))
    for label, value in totals:
        out.append(f"{label:>{width - 14}} {value:>13,.2f}")
    out.append(rule)
    out.append(f"Paid by: {sales_service.payment_label(sale)}"
               + (f"   Received: {sale.tendered_amount:.2f}   Change: {sale.change_amount or 0:.2f}" if sale.tendered_amount is not None else ""))
    return "\n".join(out) + "\n"
