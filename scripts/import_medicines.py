"""Import a real medicine catalogue (CSV or Excel) into the item master.

The India medicine database (300k+ items) is supplied by the pharmacy; this
script loads it. Column names are matched case-insensitively and a set of
common aliases is accepted.

Usage:
    .venv/bin/python scripts/import_medicines.py --file medicines.csv
    .venv/bin/python scripts/import_medicines.py --file medicines.xlsx --limit 1000
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
import time
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

ALIASES = {
    "name": ["name", "product_name", "item_name", "medicine_name", "description", "particulars"],
    "generic_name": [
        "generic_name", "generic", "composition", "salt", "molecule",
        "short_composition1", "short_composition2",
    ],
    "manufacturer": ["manufacturer", "manufacturer_name", "company", "mfr", "brand", "marketed_by"],
    "category": ["category", "type", "group"],
    "pack_size": ["pack_size", "pack_size_label", "pack", "packing", "package"],
    "strength": ["strength", "power", "dosage"],
    "unit": ["unit", "uom"],
    "hsn_code": ["hsn_code", "hsn"],
    "gst_rate": ["gst_rate", "gst", "tax"],
    "mrp": ["mrp", "price", "rate"],
    "barcode": ["barcode", "ean", "upc"],
}

CATEGORY_MAP = {
    "allopathy": "PHARMA",
    "ayurvedic": "AYURVEDIC",
    "homeopathy": "OTHER",
    "unani": "OTHER",
    "siddha": "OTHER",
    "generic": "GENERIC",
    "otc": "FMCG",
}


def norm_key(key: str) -> str:
    return re.sub(r"[^a-z0-9]", "", (key or "").lower())


def build_index(headers: list[str]) -> dict[str, str]:
    lookup = {norm_key(h): h for h in headers}
    mapping: dict[str, str] = {}
    for target, names in ALIASES.items():
        normalized = [norm_key(n) for n in names]
        for n in normalized:
            if n in lookup:
                mapping[target] = lookup[n]
                break
        if target in mapping:
            continue
        for header_norm, header in lookup.items():
            if any(header_norm.startswith(n) for n in normalized if n):
                mapping[target] = header
                break
    return mapping


def map_category(value: str) -> str:
    key = (value or "").strip().lower()
    return CATEGORY_MAP.get(key, "PHARMA")


def read_rows(path: Path):
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        headers = [str(c) if c is not None else "" for c in next(rows)]
        for row in rows:
            yield headers, row
    else:
        with path.open(newline="", encoding="utf-8-sig") as fh:
            reader = csv.DictReader(fh)
            headers = reader.fieldnames or []
            for row in reader:
                yield headers, [row.get(h, "") for h in headers]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--file", required=True)
    ap.add_argument("--limit", type=int, default=0)
    ap.add_argument("--batch-size", type=int, default=5000)
    ap.add_argument("--skip-existing", action="store_true", default=True)
    args = ap.parse_args()

    from sqlalchemy import text

    from app.database import SessionLocal, engine, ensure_fts, init_db, rebuild_fts
    from app.models import Item, NumberSequence
    from app.sequences import article_prefix
    from app.utils import to_decimal

    init_db()

    # Drop FTS triggers during the bulk load; they are recreated and the index
    # rebuilt at the end (much faster than firing a trigger per row).
    with engine.begin() as conn:
        for trig in ("items_fts_au", "items_fts_ad", "items_fts_ai"):
            conn.execute(text(f"DROP TRIGGER IF EXISTS {trig}"))

    path = Path(args.file)
    if not path.exists():
        print(f"File not found: {path}")
        sys.exit(1)

    db = SessionLocal()
    existing = {(i.name.strip().lower(), i.manufacturer, i.pack_size) for i in db.query(Item).all()}

    # Per-prefix serial counters, seeded from any items already present so new
    # IDs never collide with manual entries.
    counters: dict[str, int] = {}
    for (article_id,) in db.query(Item.article_id).all():
        prefix = (article_id or "")[:3]
        tail = (article_id or "")[3:]
        if prefix and tail.isdigit():
            counters[prefix] = max(counters.get(prefix, 0), int(tail))

    inserted = 0
    skipped = 0
    buffer: list[dict] = []
    mapping: dict[str, str] = {}
    t0 = time.time()
    now = datetime.utcnow()

    def next_article(name: str) -> str:
        prefix = article_prefix(name)
        counters[prefix] = counters.get(prefix, 0) + 1
        return f"{prefix}{counters[prefix]:04d}"

    def flush():
        nonlocal inserted
        if not buffer:
            return
        db.bulk_insert_mappings(Item, buffer)
        db.commit()
        inserted += len(buffer)
        buffer.clear()
        print(f"  ... {inserted:,} inserted", flush=True)

    for headers, row in read_rows(path):
        if not mapping:
            mapping = build_index(headers)
            if "name" not in mapping:
                print(f"Could not find a name column in headers: {headers}")
                sys.exit(1)
        rec = {target: row[headers.index(col)] for target, col in mapping.items() if col in headers}
        name = str(rec.get("name") or "").strip()
        if not name:
            skipped += 1
            continue
        manufacturer = str(rec.get("manufacturer") or "").strip()
        pack = str(rec.get("pack_size") or "").strip()
        key = (name.lower(), manufacturer, pack)
        if args.skip_existing and key in existing:
            skipped += 1
            continue
        existing.add(key)
        buffer.append(
            {
                "article_id": next_article(name),
                "name": name[:250],
                "generic_name": str(rec.get("generic_name") or "")[:250],
                "manufacturer": manufacturer[:150],
                "category": map_category(str(rec.get("category") or ""))[:60],
                "pack_size": pack[:60],
                "strength": str(rec.get("strength") or "")[:60],
                "unit": str(rec.get("unit") or "unit")[:20],
                "hsn_code": str(rec.get("hsn_code") or "")[:20],
                "gst_rate": to_decimal(rec.get("gst_rate") or 12),
                "mrp": to_decimal(rec.get("mrp") or 0),
                "barcode": str(rec.get("barcode") or "")[:60],
                "is_active": True,
                "created_at": now,
                "updated_at": now,
            }
        )
        if len(buffer) >= args.batch_size:
            flush()
        if args.limit and inserted + len(buffer) >= args.limit:
            break

    flush()

    # Persist per-prefix counters so manually added items continue the sequence.
    for prefix, used in counters.items():
        seq = db.get(NumberSequence, f"article:{prefix}")
        if seq is None:
            db.add(NumberSequence(key=f"article:{prefix}", next_value=used + 1))
        else:
            seq.next_value = max(seq.next_value, used + 1)
    db.commit()
    db.close()
    ensure_fts()
    rebuild_fts()
    print(f"Done. Inserted {inserted:,}, skipped {skipped:,} in {time.time() - t0:.1f}s")


if __name__ == "__main__":
    main()
